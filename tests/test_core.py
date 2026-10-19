import io
import shutil
from pathlib import Path

import numpy as np
import pymupdf
import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from filefit.core import Options, combine_to_pdf, convert_file

MB = 1024 * 1024


def make_photo(path: Path, size=(3000, 2200)):
    rng = np.random.default_rng(1)
    base = np.linspace(0, 255, size[0], dtype=np.float32)[None, :, None].repeat(size[1], 0).repeat(3, 2)
    noise = rng.normal(0, 40, (size[1], size[0], 3))
    Image.fromarray(np.clip(base + noise, 0, 255).astype("uint8")).save(path, quality=98)
    return path


@pytest.fixture()
def files(tmp_path):
    photo = make_photo(tmp_path / "photo.jpg")
    png = tmp_path / "logo.png"
    im = Image.new("RGBA", (800, 600), (0, 0, 0, 0))
    im.paste((30, 90, 200, 255), (100, 100, 700, 500))
    im.save(png)
    pdf = tmp_path / "doc.pdf"
    with pymupdf.open() as d:
        for i in range(2):
            p = d.new_page()
            p.insert_text((72, 72), f"Hello page {i + 1}", fontsize=24)
            p.insert_image(pymupdf.Rect(72, 120, 520, 700), filename=str(photo))
        d.save(pdf)
    docx = tmp_path / "letter.docx"
    doc = Document()
    doc.add_paragraph("Application letter")
    doc.add_picture(str(photo), width=Inches(5))
    doc.save(docx)
    return dict(photo=photo, png=png, pdf=pdf, docx=docx, out=tmp_path / "out")


def run(src, **kw):
    res = convert_file(src, Options(**kw))
    assert res.error is None, res.error
    assert res.outputs
    return res


@pytest.mark.parametrize("fmt", ["JPG", "JPEG", "PNG", "WEBP", "PDF", "DOCX"])
@pytest.mark.parametrize("key", ["photo", "png", "pdf"])
def test_every_conversion(files, key, fmt):
    res = run(files[key], out_format=fmt, out_dir=files["out"])
    for o in res.outputs:
        assert o.path.exists() and o.path.suffix == {"JPG": ".jpg", "JPEG": ".jpeg"}.get(fmt, "." + fmt.lower())


@pytest.mark.parametrize("fmt,limit", [("JPG", 2 * MB), ("JPG", 100 * 1024), ("PNG", 300 * 1024),
                                       ("PDF", 200 * 1024), ("DOCX", 300 * 1024)])
def test_size_limit_photo(files, fmt, limit):
    assert files["photo"].stat().st_size > limit
    res = run(files["photo"], out_format=fmt, out_dir=files["out"], target_bytes=limit)
    o = res.outputs[0]
    assert o.fits and o.size <= limit, (o.size, o.note)


def test_pdf_compress(files):
    limit = 150 * 1024
    assert files["pdf"].stat().st_size > limit
    res = run(files["pdf"], out_format="PDF", out_dir=files["out"], target_bytes=limit)
    assert res.outputs[0].size <= limit


def test_docx_compress(files):
    limit = 200 * 1024
    assert files["docx"].stat().st_size > limit
    res = run(files["docx"], out_format="DOCX", out_dir=files["out"], target_bytes=limit)
    assert res.outputs[0].size <= limit
    Document(str(res.outputs[0].path))  # still opens


def test_pdf_pages_become_images(files):
    res = run(files["pdf"], out_format="JPG", out_dir=files["out"])
    assert [o.path.name for o in res.outputs] == ["doc_page1.jpg", "doc_page2.jpg"]


def test_resize_dimensions(files):
    res = run(files["photo"], out_format="PNG", out_dir=files["out"], max_width=600, max_height=600)
    assert max(Image.open(res.outputs[0].path).size) == 600


def test_never_overwrites(files):
    res = run(files["photo"], out_format="JPG")
    assert res.outputs[0].path != files["photo"]


def test_combine(files):
    res = combine_to_pdf([files["photo"], files["png"], files["pdf"]],
                         Options(out_format="PDF", out_dir=files["out"], target_bytes=MB))
    assert res.error is None
    with pymupdf.open(res.outputs[0].path) as d:
        assert d.page_count == 4
    assert res.outputs[0].size <= MB


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="needs LibreOffice")
def test_docx_to_pdf_and_jpg(files):
    run(files["docx"], out_format="PDF", out_dir=files["out"])
    run(files["docx"], out_format="JPG", out_dir=files["out"], target_bytes=200 * 1024)
