"""Conversion engine: turns any supported input into any supported output.

The engine has no GUI code in it, so the same logic can be reused by a
command line tool, a test suite, or a future mobile front end.
"""
from __future__ import annotations

import io
import os
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable

import pymupdf
from PIL import Image, ImageOps

from .compress import FitResult, encode_image, fit_docx, fit_image, fit_pdf
from .formats import IMAGE_OUTPUTS, OUTPUT_FORMATS, kind_of

try:  # iPhone photos (HEIC); optional
    from pillow_heif import register_heif_opener

    register_heif_opener()
except Exception:  # pragma: no cover
    pass

Log = Callable[[str], None]

A4 = (595.0, 842.0)  # points
PAGE_MARGIN = 18.0


class ConversionError(Exception):
    pass


@dataclass
class Options:
    out_format: str = "PDF"              # a key of OUTPUT_FORMATS
    out_dir: Path | None = None          # None = next to the original file
    target_bytes: int | None = None      # size limit, or None
    max_width: int | None = None         # resize limits in pixels (images only)
    max_height: int | None = None
    combine: bool = False                # merge every input into one PDF
    render_dpi: int = 150                # resolution when turning pages into images


@dataclass
class Output:
    path: Path
    size: int
    fits: bool = True
    note: str = ""
    source: Path | None = None


@dataclass
class JobResult:
    source: Path
    outputs: list[Output] = field(default_factory=list)
    error: str | None = None


# ---------------------------------------------------------------- helpers

def unique_path(folder: Path, stem: str, ext: str, avoid: Path | None = None) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    cand = folder / f"{stem}{ext}"
    n = 1
    while cand.exists() or (avoid is not None and cand.resolve() == avoid.resolve()):
        cand = folder / f"{stem} ({n}){ext}"
        n += 1
    return cand


def limit_dims(img: Image.Image, max_w: int | None, max_h: int | None) -> Image.Image:
    if not max_w and not max_h:
        return img
    w, h = img.size
    factor = min((max_w or w) / w, (max_h or h) / h, 1.0)
    if factor >= 1.0:
        return img
    return img.resize((max(1, round(w * factor)), max(1, round(h * factor))), Image.LANCZOS)


def open_image(path: Path) -> Image.Image:
    img = Image.open(path)
    try:
        img.seek(0)
    except EOFError:
        pass
    img = ImageOps.exif_transpose(img)  # respect phone camera rotation
    img.load()
    return img


def pdf_pages_to_images(pdf_bytes: bytes, dpi: int) -> list[Image.Image]:
    images = []
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ConversionError("This PDF is password protected. Remove the password first.")
        for page in doc:
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            images.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
    return images


def _find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for base in (os.environ.get("PROGRAMFILES", r"C:\Program Files"),
                 os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)")):
        cand = Path(base) / "LibreOffice" / "program" / "soffice.exe"
        if cand.exists():
            return str(cand)
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def docx_to_pdf_bytes(src: Path, log: Log) -> bytes:
    """Convert DOCX to PDF using Microsoft Word (Windows) or LibreOffice."""
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / (src.stem + ".pdf")
        if sys.platform == "win32":
            try:
                import pythoncom
                import win32com.client

                log("using Microsoft Word to read the document")
                pythoncom.CoInitialize()
                word = win32com.client.DispatchEx("Word.Application")
                word.Visible = False
                word.DisplayAlerts = 0
                try:
                    doc = word.Documents.Open(str(src.resolve()), ReadOnly=True)
                    doc.SaveAs(str(out), FileFormat=17)  # 17 = PDF
                    doc.Close(False)
                finally:
                    word.Quit()
                if out.exists():
                    return out.read_bytes()
            except Exception as exc:
                log(f"Word not available ({exc.__class__.__name__}), trying LibreOffice")
        soffice = _find_soffice()
        if soffice:
            log("using LibreOffice to read the document")
            flags = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                           check=False, capture_output=True, timeout=180, creationflags=flags)
            if out.exists():
                return out.read_bytes()
        raise ConversionError(
            "Converting DOCX needs Microsoft Word or the free LibreOffice installed on this PC.")


def images_to_pdf(images: list[Image.Image], target: int | None) -> FitResult:
    """Put each image on its own A4 page, sized to fit, within the byte budget."""
    budget = None
    if target:
        budget = max(8_000, (target - 4_000 - 1_500 * len(images)) // max(1, len(images)))
    fits, notes = True, set()
    with pymupdf.open() as doc:
        for img in images:
            res = fit_image(img, "JPEG", budget, quality=90)
            fits &= res.fits
            if res.note:
                notes.add(res.note)
            w, h = img.size
            pw, ph = A4 if h >= w else (A4[1], A4[0])
            page = doc.new_page(width=pw, height=ph)
            box = pymupdf.Rect(PAGE_MARGIN, PAGE_MARGIN, pw - PAGE_MARGIN, ph - PAGE_MARGIN)
            page.insert_image(box, stream=res.data, keep_proportion=True)
        data = doc.tobytes(garbage=4, deflate=True)
    fits = target is None or len(data) <= target
    return FitResult(data, fits, "; ".join(sorted(notes)))


def images_to_docx(images: list[Image.Image], target: int | None) -> FitResult:
    from docx import Document
    from docx.shared import Inches

    empty = io.BytesIO()
    Document().save(empty)
    overhead = len(empty.getvalue()) + 2_000
    budget = None
    if target:
        budget = max(8_000, (target - overhead) // max(1, len(images)))
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.7)
        section.top_margin = section.bottom_margin = Inches(0.7)
    usable_w, usable_h = 8.27 - 1.4, 11.69 - 1.4  # A4 in inches
    notes = set()
    for i, img in enumerate(images):
        res = fit_image(img, "JPEG", budget, quality=90)
        if res.note:
            notes.add(res.note)
        w, h = img.size
        scale = min(usable_w / w, usable_h / h)
        if i:
            doc.add_page_break()
        doc.add_picture(io.BytesIO(res.data), width=Inches(w * scale))
    out = io.BytesIO()
    doc.save(out)
    data = out.getvalue()
    return FitResult(data, target is None or len(data) <= target, "; ".join(sorted(notes)))


def pdf_to_docx_bytes(pdf_bytes: bytes, log: Log, dpi: int) -> tuple[bytes, str]:
    """Editable conversion with pdf2docx, falling back to page pictures."""
    try:
        from pdf2docx import Converter

        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "in.pdf"
            dst = Path(tmp) / "out.docx"
            src.write_bytes(pdf_bytes)
            log("rebuilding the PDF as an editable Word document")
            cv = Converter(str(src))
            try:
                cv.convert(str(dst))
            finally:
                cv.close()
            if dst.exists() and dst.stat().st_size > 0:
                return dst.read_bytes(), ""
    except Exception as exc:
        log(f"editable conversion failed ({exc}); using page pictures instead")
    res = images_to_docx(pdf_pages_to_images(pdf_bytes, dpi), None)
    return res.data, "pages inserted as pictures (not editable)"


# ---------------------------------------------------------------- main API

def _load_as_pdf(src: Path, log: Log) -> bytes:
    kind = kind_of(src)
    if kind == "pdf":
        return src.read_bytes()
    if kind == "docx":
        return docx_to_pdf_bytes(src, log)
    return images_to_pdf([open_image(src)], None).data


def convert_file(src: Path, opts: Options, log: Log = lambda s: None) -> JobResult:
    src = Path(src)
    result = JobResult(source=src)
    try:
        ext, fmt = OUTPUT_FORMATS[opts.out_format]
        kind = kind_of(src)
        folder = Path(opts.out_dir) if opts.out_dir else src.parent
        target = opts.target_bytes

        def write(stem: str, res: FitResult) -> None:
            if folder.resolve() == src.parent.resolve() and ext == src.suffix.lower():
                stem += "_compressed" if target else "_converted"
            path = unique_path(folder, stem, ext, avoid=src)
            path.write_bytes(res.data)
            result.outputs.append(Output(path, len(res.data), res.fits, res.note, src))

        if opts.out_format in IMAGE_OUTPUTS:
            if kind == "image":
                images = [open_image(src)]
            else:
                pdf = _load_as_pdf(src, log)
                log("turning pages into images")
                images = pdf_pages_to_images(pdf, opts.render_dpi)
            for i, img in enumerate(images, 1):
                img = limit_dims(img, opts.max_width, opts.max_height)
                stem = src.stem if len(images) == 1 else f"{src.stem}_page{i}"
                write(stem, fit_image(img, fmt, target))

        elif opts.out_format == "PDF":
            if kind == "image":
                img = limit_dims(open_image(src), opts.max_width, opts.max_height)
                write(src.stem, images_to_pdf([img], target))
            else:
                pdf = _load_as_pdf(src, log)
                if target:
                    log("compressing PDF")
                write(src.stem, fit_pdf(pdf, target, log))

        elif opts.out_format == "DOCX":
            if kind == "image":
                img = limit_dims(open_image(src), opts.max_width, opts.max_height)
                write(src.stem, images_to_docx([img], target))
            elif kind == "pdf":
                data, note = pdf_to_docx_bytes(src.read_bytes(), log, opts.render_dpi)
                res = fit_docx(data, target)
                res.note = "; ".join(n for n in (note, res.note) if n)
                write(src.stem, res)
            else:
                write(src.stem, fit_docx(src.read_bytes(), target))
        else:
            raise ConversionError(f"Unknown output format {opts.out_format}")
    except ConversionError as exc:
        result.error = str(exc)
    except Exception as exc:  # keep the app alive on any bad file
        result.error = f"{exc.__class__.__name__}: {exc}"
    return result


def combine_to_pdf(sources: Iterable[Path], opts: Options, log: Log = lambda s: None) -> JobResult:
    """Merge every input (images, PDFs, DOCX) into one PDF, in order."""
    sources = [Path(s) for s in sources]
    result = JobResult(source=sources[0])
    try:
        with pymupdf.open() as merged:
            for src in sources:
                log(f"adding {src.name}")
                if kind_of(src) == "image":
                    img = limit_dims(open_image(src), opts.max_width, opts.max_height)
                    part = images_to_pdf([img], None).data
                else:
                    part = _load_as_pdf(src, log)
                with pymupdf.open(stream=part, filetype="pdf") as d:
                    merged.insert_pdf(d)
            raw = merged.tobytes(garbage=4, deflate=True)
        res = fit_pdf(raw, opts.target_bytes, log)
        folder = Path(opts.out_dir) if opts.out_dir else sources[0].parent
        path = unique_path(folder, f"{sources[0].stem}_combined", ".pdf")
        path.write_bytes(res.data)
        result.outputs.append(Output(path, len(res.data), res.fits, res.note, sources[0]))
    except Exception as exc:
        result.error = str(exc)
    return result
